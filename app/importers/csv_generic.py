import csv
import io
from datetime import datetime
from typing import Optional

from openpyxl import load_workbook

from .base import NormalizedTransaction

DATE_FORMATS = ["%m/%d/%Y", "%Y-%m-%d", "%m/%d/%y", "%d/%m/%Y", "%B %d, %Y", "%b %d, %Y"]


def parse_date(value: str) -> Optional[str]:
    value = (value or "").strip()
    if not value:
        return None
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def parse_amount(value: str) -> Optional[float]:
    value = (value or "").strip()
    if not value:
        return None
    negative = value.startswith("(") and value.endswith(")")
    cleaned = value.replace("(", "").replace(")", "").replace("$", "").replace(",", "").strip()
    if not cleaned or cleaned in {"-", "."}:
        return None
    try:
        amount = float(cleaned)
    except ValueError:
        return None
    return -amount if negative else amount


def read_tabular(file_bytes: bytes, filename: str) -> tuple[list[str], list[list[str]]]:
    """Reads a CSV or XLSX file into (headers, rows-of-strings)."""
    lower = filename.lower()
    if lower.endswith(".xlsx") or lower.endswith(".xlsm"):
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            return [], []
        headers = [str(h).strip() if h is not None else "" for h in header_row]
        rows = []
        for row in rows_iter:
            rows.append(["" if v is None else str(v) for v in row])
        return headers, rows
    else:
        text = file_bytes.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text))
        all_rows = [r for r in reader if any(cell.strip() for cell in r)]
        if not all_rows:
            return [], []
        headers = [h.strip() for h in all_rows[0]]
        return headers, all_rows[1:]


def sniff_headers(file_bytes: bytes, filename: str) -> tuple[list[str], list[list[str]]]:
    """Returns headers and up to 5 sample rows, for the column-mapping UI."""
    headers, rows = read_tabular(file_bytes, filename)
    return headers, rows[:5]


def apply_mapping(
    headers: list[str],
    rows: list[list[str]],
    mapping: dict,
) -> list[NormalizedTransaction]:
    """
    mapping keys:
      date_col (required), posted_date_col (optional), description_col (required)
      amount_col + amount_sign ("spend_negative" | "spend_positive")  -- OR --
      debit_col + credit_col (debit = money out, credit = money in)
    """
    def col_index(name: Optional[str]) -> Optional[int]:
        if not name:
            return None
        try:
            return headers.index(name)
        except ValueError:
            return None

    date_i = col_index(mapping.get("date_col"))
    posted_i = col_index(mapping.get("posted_date_col"))
    desc_i = col_index(mapping.get("description_col"))
    amount_i = col_index(mapping.get("amount_col"))
    debit_i = col_index(mapping.get("debit_col"))
    credit_i = col_index(mapping.get("credit_col"))
    amount_sign = mapping.get("amount_sign", "spend_negative")

    def cell(row: list[str], idx: Optional[int]) -> str:
        if idx is None or idx >= len(row):
            return ""
        return row[idx] or ""

    results: list[NormalizedTransaction] = []
    for row in rows:
        date_val = parse_date(cell(row, date_i))
        if date_val is None:
            continue
        raw_desc = cell(row, desc_i).strip()
        if not raw_desc:
            continue

        amount: Optional[float] = None
        if amount_i is not None:
            amount = parse_amount(cell(row, amount_i))
            if amount is not None and amount_sign == "spend_positive":
                amount = -amount
        elif debit_i is not None or credit_i is not None:
            debit = parse_amount(cell(row, debit_i)) or 0.0
            credit = parse_amount(cell(row, credit_i)) or 0.0
            amount = credit - abs(debit)

        if amount is None:
            continue

        posted_val = parse_date(cell(row, posted_i)) if posted_i is not None else None

        results.append(
            NormalizedTransaction(
                date=date_val,
                posted_date=posted_val,
                description=" ".join(raw_desc.split()),
                raw_description=raw_desc,
                amount=amount,
            )
        )
    return results
